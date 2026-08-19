#!/usr/bin/env python3
"""2026 Road Restriction Snapshot Summary.

Plan: Report observation availability, unique records, features, reporting states,
observed active length, and active-length missingness for every official snapshot.
Framework: AnaSOP Sections 5.1-5.4, 6.1, and 7; implements N_st, N_t,
D_t, L_t^(active,obs), q_t^active, and the explicit unavailable-layer rule.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RESTRICTIONS_PATH = (
    ROOT
    / "data/processed/kumamoto_2026_road_restriction_snapshots_preprocessed.parquet"
)
COVERAGE_PATH = (
    ROOT / "data/processed/kumamoto_2026_road_snapshot_coverage_preprocessed.parquet"
)
OUTPUT_PATH = (
    ROOT / "data/exp/tables/Table_2026_road_restriction_snapshot_summary.xlsx"
)

ACTIVE_STATES = {"active_full_closure", "active_emergency_access_only"}
OUTPUT_COLUMNS = [
    "Snapshot Time",
    "Restriction Layer Available",
    "Unique Restriction Records",
    "Restriction Features",
    "Active Features",
    "Lifted Features",
    "Unknown Status Features",
    "Observed Active Length (km)",
    "Active Length Missing Share",
]


def build_summary(restrictions: pd.DataFrame, coverage: pd.DataFrame) -> pd.DataFrame:
    """Construct one row per official snapshot while preserving unavailable coverage."""
    required_restriction_columns = {
        "Snapshot Time",
        "Restriction Record ID",
        "Restriction Feature ID",
        "Observation State",
        "Restriction Length (km)",
    }
    required_coverage_columns = {
        "Snapshot Time",
        "Restriction Layer Available",
        "Restriction Feature Count",
        "Absence Interpretation",
    }
    missing_restriction = required_restriction_columns - set(restrictions.columns)
    missing_coverage = required_coverage_columns - set(coverage.columns)
    if missing_restriction or missing_coverage:
        raise ValueError(
            "Missing required columns: "
            f"restrictions={sorted(missing_restriction)}, "
            f"coverage={sorted(missing_coverage)}"
        )

    if restrictions["Restriction Feature ID"].duplicated().any():
        raise ValueError("Restriction Feature ID must be unique in the processed panel.")
    if coverage["Snapshot Time"].duplicated().any():
        raise ValueError("Coverage must contain one row per official snapshot.")

    work = restrictions.copy()
    work["Is Active"] = work["Observation State"].isin(ACTIVE_STATES)
    work["Is Lifted"] = work["Observation State"].eq("restriction_lifted")
    work["Is Unknown"] = work["Observation State"].eq("unknown_status")
    work["Has Active Length"] = work["Is Active"] & work[
        "Restriction Length (km)"
    ].notna()
    work["Observed Active Length Contribution"] = work[
        "Restriction Length (km)"
    ].where(work["Is Active"] & work["Restriction Length (km)"].notna())

    grouped = work.groupby("Snapshot Time", observed=True, sort=True)
    observed = grouped.agg(
        **{
            "Unique Restriction Records": (
                "Restriction Record ID",
                lambda s: s.dropna().nunique(),
            ),
            "Restriction Features": ("Restriction Feature ID", "nunique"),
            "Active Features": ("Is Active", "sum"),
            "Lifted Features": ("Is Lifted", "sum"),
            "Unknown Status Features": ("Is Unknown", "sum"),
            "Observed Active Length (km)": (
                "Observed Active Length Contribution",
                lambda s: s.sum(min_count=1),
            ),
            "Active Features With Length": ("Has Active Length", "sum"),
        }
    ).reset_index()

    observed["Active Length Missing Share"] = (
        1
        - observed["Active Features With Length"]
        / observed["Active Features"].replace(0, pd.NA)
    )

    summary = coverage[
        [
            "Snapshot Time",
            "Restriction Layer Available",
            "Restriction Feature Count",
            "Absence Interpretation",
        ]
    ].merge(observed, on="Snapshot Time", how="left", validate="one_to_one")

    available = summary["Restriction Layer Available"]
    metrics = OUTPUT_COLUMNS[2:]
    summary.loc[~available, metrics] = pd.NA

    available_rows = summary.loc[available]
    if not (
        available_rows["Restriction Features"].astype("Int64")
        == available_rows["Restriction Feature Count"].astype("Int64")
    ).all():
        raise ValueError("Calculated feature counts do not match the coverage audit.")
    if not (
        available_rows["Restriction Features"].astype("Int64")
        == available_rows[
            ["Active Features", "Lifted Features", "Unknown Status Features"]
        ]
        .sum(axis=1)
        .astype("Int64")
    ).all():
        raise ValueError("Observation-state counts do not sum to feature counts.")
    if not summary.loc[~available, "Absence Interpretation"].eq(
        "not_observed_not_open"
    ).all():
        raise ValueError("Unavailable layers must retain the not-observed-not-open rule.")

    summary = summary.sort_values("Snapshot Time").reset_index(drop=True)
    summary["Snapshot Time"] = summary["Snapshot Time"].dt.tz_convert(
        "Asia/Tokyo"
    ).dt.tz_localize(None)

    for column in OUTPUT_COLUMNS[2:7]:
        summary[column] = summary[column].astype("Int64")
    return summary[OUTPUT_COLUMNS]


def write_workbook(table: pd.DataFrame, output_path: Path) -> None:
    """Write the authoritative, reproducible workbook and apply publication styling."""
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Snapshot Summary", index=False)
        workbook = writer.book
        sheet = writer.sheets["Snapshot Summary"]

        notes = workbook.create_sheet("Notes")
        notes_rows = [
            ("Field", "Definition / interpretation"),
            ("Snapshot Time", "Official observation time in Japan Standard Time (JST)."),
            (
                "Restriction Layer Available",
                "FALSE means the archive supplied no restriction layer; all summary metrics "
                "remain blank and must not be interpreted as zero restrictions or open roads.",
            ),
            (
                "Unique Restriction Records",
                "Number of unique Restriction Record ID values in the available snapshot (D_t).",
            ),
            (
                "Restriction Features",
                "Number of unique Restriction Feature ID values in the available snapshot (N_t).",
            ),
            (
                "Active Features",
                "Features coded active_full_closure or active_emergency_access_only.",
            ),
            (
                "Lifted Features",
                "Features coded restriction_lifted; not counted as active.",
            ),
            (
                "Unknown Status Features",
                "Features coded unknown_status; not counted as active.",
            ),
            (
                "Observed Active Length (km)",
                "Sum of nonmissing reported lengths among active features only; not total "
                "restricted network length.",
            ),
            (
                "Active Length Missing Share",
                "Share of active features with missing Restriction Length (km) (q_t^active).",
            ),
            (
                "Interpretation boundary",
                "Counts and lengths describe the official reporting archive. A decline is not "
                "evidence of reopening or recovery without stable road-state linkage.",
            ),
            (
                "Source",
                "Processed official snapshot and coverage files listed in the generating script.",
            ),
        ]
        for row in notes_rows:
            notes.append(row)

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="1F2937")
        stripe_fill = PatternFill("solid", fgColor="F2F6FA")
        unavailable_fill = PatternFill("solid", fgColor="FCE8E6")
        note_fill = PatternFill("solid", fgColor="EAF2F8")
        thin_gray = Side(style="thin", color="D5DCE4")

        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"
        sheet.sheet_view.showGridLines = False
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.orientation = "landscape"
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.print_title_rows = "1:1"
        sheet.sheet_properties.outlinePr.summaryBelow = True

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            cell.border = Border(bottom=Side(style="medium", color="17365D"))
        sheet.row_dimensions[1].height = 44

        for row_idx in range(2, sheet.max_row + 1):
            is_available = bool(sheet.cell(row_idx, 2).value)
            row_fill = (
                unavailable_fill
                if not is_available
                else (stripe_fill if row_idx % 2 == 0 else None)
            )
            for col_idx in range(1, 10):
                cell = sheet.cell(row_idx, col_idx)
                cell.font = body_font
                cell.alignment = Alignment(
                    horizontal="left" if col_idx == 1 else "right",
                    vertical="center",
                )
                cell.border = Border(bottom=thin_gray)
                if row_fill is not None:
                    cell.fill = row_fill
            sheet.cell(row_idx, 2).alignment = Alignment(horizontal="center")

        for cell in sheet["A"][1:]:
            cell.number_format = "yyyy-mm-dd hh:mm"
        for col_idx in range(3, 8):
            for cell in sheet[get_column_letter(col_idx)][1:]:
                cell.number_format = "0"
        for cell in sheet["H"][1:]:
            cell.number_format = "0.00"
        for cell in sheet["I"][1:]:
            cell.number_format = "0.0%"

        widths = [19, 18, 18, 17, 14, 14, 18, 20, 20]
        for col_idx, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(col_idx)].width = width
        sheet.auto_filter.ref = f"A1:I{sheet.max_row}"

        sheet["B1"].comment = Comment(
            "FALSE means not observed, not zero restrictions or open roads.",
            "MiliFrame",
        )
        sheet["H1"].comment = Comment(
            "Sum among active features with nonmissing reported length only.",
            "MiliFrame",
        )
        sheet["I1"].comment = Comment(
            "Missing active-length fields divided by all active features.",
            "MiliFrame",
        )

        notes.freeze_panes = "A2"
        notes.sheet_view.showGridLines = False
        notes.column_dimensions["A"].width = 34
        notes.column_dimensions["B"].width = 104
        notes.row_dimensions[1].height = 28
        for cell in notes[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="medium", color="17365D"))
        for row_idx in range(2, notes.max_row + 1):
            notes.cell(row_idx, 1).font = Font(
                name="Aptos", size=10, bold=True, color="17365D"
            )
            notes.cell(row_idx, 2).font = body_font
            for col_idx in (1, 2):
                cell = notes.cell(row_idx, col_idx)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin_gray)
                if row_idx % 2 == 0:
                    cell.fill = note_fill
            notes.row_dimensions[row_idx].height = 34


def main() -> None:
    restrictions = pd.read_parquet(RESTRICTIONS_PATH)
    coverage = pd.read_parquet(COVERAGE_PATH)
    table = build_summary(restrictions, coverage)
    if table.shape != (24, 9):
        raise ValueError(f"Expected a 24 x 9 table, obtained {table.shape}.")
    write_workbook(table, OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH.relative_to(ROOT)}")
    print(
        "Checks: 24 official snapshots; "
        f"{int(table['Restriction Layer Available'].sum())} available; "
        f"{int((~table['Restriction Layer Available']).sum())} unavailable; "
        "state counts reconcile with feature counts."
    )


if __name__ == "__main__":
    main()
