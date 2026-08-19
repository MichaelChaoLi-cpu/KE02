#!/usr/bin/env python3
"""Reported Restriction Composition and Persistence.

Plan: Summarize record-weighted category composition, feature-snapshot
sensitivity, reporting persistence, continuity, observation bounds, and shares
renormalized after excluding Not reported.
Framework: AnaSOP Sections 5.1-5.4, 6.1, and 7; implements P_j, C_j,
rho_c^record, rho_c^feature, rho_c^record_nm, and rho_c^feature_nm.
Persistence is not closure duration.
"""
from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


ROOT = Path(__file__).resolve().parents[2]
RESTRICTIONS_PATH = (
    ROOT
    / "data/processed/kumamoto_2026_road_restriction_snapshots_preprocessed.parquet"
)
COVERAGE_PATH = (
    ROOT / "data/processed/kumamoto_2026_road_snapshot_coverage_preprocessed.parquet"
)
OUTPUT_PATH = (
    ROOT
    / "docs/process-records/assets/supplementary/Table_reported_restriction_composition_and_persistence.xlsx"
)

OUTPUT_COLUMNS = [
    "Category Family",
    "Category",
    "Unique Record Count",
    "Unique Record Share",
    "Feature-Snapshot Share",
    "Median Reported Snapshots",
    "Gap-Free Share",
    "First Observed Time",
    "Last Observed Time",
]

CATEGORY_CONFIG = [
    ("Observation State", "Observation state"),
    ("Road Type", "Road type"),
    ("Restriction Reason", "Restriction reason"),
]

CATEGORY_LABELS = {
    "active_full_closure": "Active full closure",
    "active_emergency_access_only": "Active, emergency access only",
    "restriction_lifted": "Restriction lifted",
    "unknown_status": "Unknown status",
    "一般国道": "General national road",
    "国道445号": "National Route 445",
    "市区町村道": "Municipal road",
    "有料道路": "Toll road",
    "直轄国道": "Directly managed national road",
    "補助国道": "Subsidized national road",
    "都道府県道": "Prefectural road",
    "高速道路": "Expressway",
    "落石": "Rockfall",
    "橋梁損傷": "Bridge damage",
    "道路損壊": "Road damage",
    "地震": "Earthquake",
    "落下物・飛来物": "Falling or flying objects",
    "法面崩落": "Slope failure",
    "地震による路面段差": "Earthquake-related surface step",
    "土砂流入": "Sediment inflow",
    "倒木": "Fallen trees",
    "地震による橋台段差": "Earthquake-related abutment step",
    "Missing": "Not reported",
}


def canonical_nonmissing(series: pd.Series) -> tuple[str, bool]:
    """Return deterministic modal nonmissing category and inconsistency flag."""
    nonmissing = series.dropna().astype("string")
    if nonmissing.empty:
        return "Missing", False
    counts = nonmissing.value_counts()
    maximum = counts.max()
    tied_modes = sorted(str(value) for value in counts[counts.eq(maximum)].index)
    return tied_modes[0], nonmissing.nunique() > 1


def build_record_persistence(
    restrictions: pd.DataFrame, coverage: pd.DataFrame
) -> pd.DataFrame:
    available_times = list(
        coverage.loc[
            coverage["Restriction Layer Available"], "Snapshot Time"
        ].sort_values()
    )
    if len(available_times) != 24:
        raise ValueError("Expected 24 snapshots with an available restriction layer.")
    positions = {pd.Timestamp(time): index for index, time in enumerate(available_times)}

    rows: list[dict[str, object]] = []
    for record_id, group in restrictions.groupby("Restriction Record ID", sort=True):
        observed_times = sorted(
            pd.Timestamp(time) for time in group["Snapshot Time"].unique()
        )
        observed_positions = [positions[time] for time in observed_times]
        expected_count = observed_positions[-1] - observed_positions[0] + 1
        row: dict[str, object] = {
            "Restriction Record ID": record_id,
            "Observed Snapshot Count": len(observed_times),
            "Continuity Ratio": len(observed_times) / expected_count,
            "Gap Free": len(observed_times) == expected_count,
            "First Record Observation": observed_times[0],
            "Last Record Observation": observed_times[-1],
        }
        for source_column, _ in CATEGORY_CONFIG:
            canonical, inconsistent = canonical_nonmissing(group[source_column])
            row[f"Canonical {source_column}"] = canonical
            row[f"Inconsistent {source_column}"] = inconsistent
        rows.append(row)

    persistence = pd.DataFrame(rows)
    if len(persistence) != 163:
        raise ValueError(f"Expected 163 unique records, obtained {len(persistence)}.")
    if int(persistence["Gap Free"].sum()) != 134:
        raise ValueError("Expected 134 records without an internal reporting gap.")
    return persistence


def build_summary(
    restrictions: pd.DataFrame, persistence: pd.DataFrame
) -> tuple[pd.DataFrame, dict[str, int]]:
    total_records = persistence["Restriction Record ID"].nunique()
    total_feature_snapshots = restrictions["Restriction Feature ID"].nunique()
    if total_feature_snapshots != len(restrictions):
        raise ValueError("Restriction Feature ID must be unique in the processed panel.")

    rows: list[dict[str, object]] = []
    inconsistent_counts: dict[str, int] = {}

    for source_column, family_label in CATEGORY_CONFIG:
        canonical_column = f"Canonical {source_column}"
        inconsistent_column = f"Inconsistent {source_column}"
        inconsistent_counts[family_label] = int(persistence[inconsistent_column].sum())

        feature_category = (
            restrictions[source_column].fillna("Missing").astype("string")
        )
        categories = sorted(
            set(feature_category.astype(str))
            | set(persistence[canonical_column].astype(str))
        )

        family_rows: list[dict[str, object]] = []
        for category in categories:
            record_mask = persistence[canonical_column].eq(category)
            feature_mask = feature_category.eq(category)
            record_count = int(record_mask.sum())
            category_records = persistence.loc[record_mask]
            category_features = restrictions.loc[feature_mask]
            family_rows.append(
                {
                    "Category Family": family_label,
                    "Category": CATEGORY_LABELS.get(category, category),
                    "Unique Record Count": record_count,
                    "Unique Record Share": record_count / total_records,
                    "Feature-Snapshot Share": int(feature_mask.sum())
                    / total_feature_snapshots,
                    "Median Reported Snapshots": (
                        float(category_records["Observed Snapshot Count"].median())
                        if record_count
                        else np.nan
                    ),
                    "Gap-Free Share": (
                        float(category_records["Gap Free"].mean())
                        if record_count
                        else np.nan
                    ),
                    "First Observed Time": (
                        category_features["Snapshot Time"].min()
                        if not category_features.empty
                        else pd.NaT
                    ),
                    "Last Observed Time": (
                        category_features["Snapshot Time"].max()
                        if not category_features.empty
                        else pd.NaT
                    ),
                }
            )

        family = pd.DataFrame(family_rows).sort_values(
            ["Unique Record Share", "Feature-Snapshot Share", "Category"],
            ascending=[False, False, True],
        )
        if not np.isclose(family["Unique Record Share"].sum(), 1):
            raise ValueError(f"Record shares do not sum to one for {family_label}.")
        if not np.isclose(family["Feature-Snapshot Share"].sum(), 1):
            raise ValueError(
                f"Feature-snapshot shares do not sum to one for {family_label}."
            )
        rows.extend(family.to_dict("records"))

    summary = pd.DataFrame(rows, columns=OUTPUT_COLUMNS)
    for column in ("First Observed Time", "Last Observed Time"):
        summary[column] = pd.to_datetime(summary[column], utc=True).dt.tz_convert(
            "Asia/Tokyo"
        ).dt.tz_localize(None)
    summary["Unique Record Count"] = summary["Unique Record Count"].astype("Int64")
    return summary, inconsistent_counts


def build_missing_excluded_summary(table: pd.DataFrame) -> pd.DataFrame:
    """Remove Not reported and renormalize both composition shares by family."""
    sensitivity = table.loc[table["Category"].ne("Not reported")].copy()
    share_columns = ["Unique Record Share", "Feature-Snapshot Share"]
    for family, family_rows in sensitivity.groupby("Category Family", sort=False):
        for column in share_columns:
            denominator = float(family_rows[column].sum())
            if denominator <= 0:
                raise ValueError(
                    f"Cannot renormalize {column} for category family {family}."
                )
            sensitivity.loc[family_rows.index, column] = (
                family_rows[column] / denominator
            )
    for family, family_rows in sensitivity.groupby("Category Family", sort=False):
        for column in share_columns:
            if not np.isclose(family_rows[column].sum(), 1):
                raise ValueError(
                    f"Missing-excluded {column} does not sum to one for {family}."
                )
    return sensitivity.reset_index(drop=True)


def write_workbook(
    table: pd.DataFrame,
    missing_excluded: pd.DataFrame,
    inconsistent_counts: dict[str, int],
    output_path: Path,
) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(output_path, engine="openpyxl") as writer:
        table.to_excel(writer, sheet_name="Composition & Persistence", index=False)
        missing_excluded.to_excel(
            writer, sheet_name="Excluding Not Reported", index=False
        )
        workbook = writer.book
        sheet = writer.sheets["Composition & Persistence"]
        sensitivity_sheet = writer.sheets["Excluding Not Reported"]

        notes = workbook.create_sheet("Notes")
        notes_rows = [
            ("Field", "Definition / interpretation"),
            (
                "Workbook sheets",
                "Composition & Persistence retains the explicit Not reported category. "
                "Excluding Not Reported removes that category and renormalizes Unique "
                "Record Share and Feature-Snapshot Share to 100% within each family.",
            ),
            (
                "Unique Record Count / Share",
                "Number and share of 163 unique Restriction Record ID values assigned "
                "to the modal nonmissing category within each family.",
            ),
            (
                "Feature-Snapshot Share",
                "Share of 1,109 Restriction Feature ID rows. This weighting gives more "
                "influence to records appearing in more snapshots or multiple features.",
            ),
            (
                "Median Reported Snapshots",
                "Median number of available snapshots in which records assigned to the "
                "category appear (P_j). This is reporting persistence, not duration.",
            ),
            (
                "Gap-Free Share",
                "Share of category-assigned records observed in every available snapshot "
                "between first and last appearance (C_j = 1).",
            ),
            (
                "First / Last Observed Time",
                "Earliest and latest Japan Standard Time at which the source category "
                "appears in a feature-snapshot row; these are observation bounds, not "
                "closure or restoration times.",
            ),
            (
                "Missing categories",
                "Source missing values remain an explicit Not reported category in the "
                "primary sheet and are excluded only in the companion sensitivity sheet.",
            ),
            (
                "Category consistency",
                "Records with multiple nonmissing category values are retained and flagged "
                "during construction; the displayed category is the deterministic modal "
                "nonmissing value.",
            ),
            (
                "Inconsistent records",
                "; ".join(
                    f"{family}: {count}" for family, count in inconsistent_counts.items()
                ),
            ),
            (
                "Interpretation boundary",
                "The table describes composition and persistence in the official reporting "
                "archive. It does not estimate restriction prevalence, closure duration, "
                "reopening, restoration, or causal resilience effects.",
            ),
            (
                "Source",
                "Processed official restriction snapshots and snapshot-coverage records "
                "listed in the generating script.",
            ),
        ]
        for row in notes_rows:
            notes.append(row)

        header_fill = PatternFill("solid", fgColor="17365D")
        header_font = Font(name="Aptos", size=10, bold=True, color="FFFFFF")
        body_font = Font(name="Aptos", size=10, color="1F2937")
        stripe_fill = PatternFill("solid", fgColor="F2F6FA")
        family_fill = PatternFill("solid", fgColor="E6EEF5")
        note_fill = PatternFill("solid", fgColor="EAF2F8")
        thin_gray = Side(style="thin", color="D5DCE4")
        medium_blue = Side(style="medium", color="7B9CB4")

        def style_data_sheet(data_sheet: object) -> None:
            data_sheet.freeze_panes = "C2"
            data_sheet.auto_filter.ref = f"A1:I{data_sheet.max_row}"
            data_sheet.sheet_view.showGridLines = False
            data_sheet.sheet_properties.pageSetUpPr.fitToPage = True
            data_sheet.page_setup.orientation = "landscape"
            data_sheet.page_setup.fitToWidth = 1
            data_sheet.page_setup.fitToHeight = 1
            data_sheet.print_title_rows = "1:1"
            data_sheet.print_area = f"A1:I{data_sheet.max_row}"

            for cell in data_sheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(
                    horizontal="center", vertical="center", wrap_text=True
                )
                cell.border = Border(bottom=Side(style="medium", color="17365D"))
            data_sheet.row_dimensions[1].height = 46

            previous_family: str | None = None
            for row_idx in range(2, data_sheet.max_row + 1):
                family = str(data_sheet.cell(row_idx, 1).value)
                starts_family = family != previous_family
                for col_idx in range(1, 10):
                    cell = data_sheet.cell(row_idx, col_idx)
                    cell.font = body_font
                    cell.alignment = Alignment(
                        horizontal="left" if col_idx <= 2 else "right",
                        vertical="center",
                    )
                    cell.border = Border(
                        top=medium_blue if starts_family else Side(style=None),
                        bottom=thin_gray,
                    )
                    if starts_family:
                        cell.fill = family_fill
                    elif row_idx % 2 == 0:
                        cell.fill = stripe_fill
                if starts_family:
                    data_sheet.cell(row_idx, 1).font = Font(
                        name="Aptos", size=10, bold=True, color="17365D"
                    )
                previous_family = family

            for cell in data_sheet["C"][1:]:
                cell.number_format = "#,##0"
            for col_idx in (4, 5, 7):
                for cell in data_sheet[get_column_letter(col_idx)][1:]:
                    cell.number_format = "0.0%"
            for cell in data_sheet["F"][1:]:
                cell.number_format = "0.0"
            for col_idx in (8, 9):
                for cell in data_sheet[get_column_letter(col_idx)][1:]:
                    cell.number_format = "yyyy-mm-dd hh:mm"

            widths = [20, 35, 17, 17, 19, 21, 15, 19, 19]
            for col_idx, width in enumerate(widths, start=1):
                data_sheet.column_dimensions[get_column_letter(col_idx)].width = width

        style_data_sheet(sheet)
        style_data_sheet(sensitivity_sheet)

        notes.freeze_panes = "A2"
        notes.sheet_view.showGridLines = False
        notes.sheet_properties.pageSetUpPr.fitToPage = True
        notes.page_setup.orientation = "landscape"
        notes.page_setup.fitToWidth = 1
        notes.page_setup.fitToHeight = 1
        notes.print_title_rows = "1:1"
        notes.print_area = f"A1:B{notes.max_row}"
        notes.column_dimensions["A"].width = 32
        notes.column_dimensions["B"].width = 106
        notes.row_dimensions[1].height = 28
        for cell in notes[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(vertical="center")
            cell.border = Border(bottom=Side(style="medium", color="17365D"))
        for row_idx in range(2, notes.max_row + 1):
            notes.cell(row_idx, 1).font = Font(
                name="Aptos", size=10, bold=True, color="17365D"
            )
            notes.cell(row_idx, 2).font = body_font
            for col_idx in (1, 2):
                cell = notes.cell(row_idx, col_idx)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.border = Border(bottom=thin_gray)
                if row_idx % 2 == 0:
                    cell.fill = note_fill
            notes.row_dimensions[row_idx].height = 36


def main() -> None:
    restrictions = pd.read_parquet(RESTRICTIONS_PATH)
    coverage = pd.read_parquet(COVERAGE_PATH)
    persistence = build_record_persistence(restrictions, coverage)
    table, inconsistent_counts = build_summary(restrictions, persistence)
    missing_excluded = build_missing_excluded_summary(table)
    if table.shape != (24, 9):
        raise ValueError(f"Expected a 24 x 9 table, obtained {table.shape}.")
    if missing_excluded.shape != (22, 9):
        raise ValueError(
            "Expected a 22 x 9 missing-excluded table, obtained "
            f"{missing_excluded.shape}."
        )
    write_workbook(table, missing_excluded, inconsistent_counts, OUTPUT_PATH)
    print(f"Saved: {OUTPUT_PATH.relative_to(ROOT)}")
    print(
        "Checks: 163 unique records; 1,109 feature-snapshot rows; "
        "134 gap-free records; family shares reconcile to 100%; "
        "missing-excluded family shares reconcile to 100%; "
        f"inconsistent records by family={inconsistent_counts}."
    )


if __name__ == "__main__":
    main()
