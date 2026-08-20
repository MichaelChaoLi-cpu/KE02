#!/usr/bin/env python3
"""Claim Robustness and Inference Ceilings.

Plan: Group twelve prespecified claim-scenario checks into four substantive
claim families and state the result, measurement dependence, and ceiling.
Framework: AnaSOP Sections 5.2-5.5, Section 6.5, and Analytical Workflow
step 6. Classifications remain categorical and cases remain non-pooled.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "data/processed"
    / "repeated_shock_claim_sensitivity_preprocessed.parquet"
)
OUT = (
    ROOT
    / "data/results/tables"
    / "Table_claim_robustness_and_inference_ceilings.xlsx"
)

NAVY = "243746"
WHITE = "FFFFFF"
TEXT = "263238"
LIGHT_GRAY = "E9EDF0"
ROW_GRAY = "F7F8F9"
ORANGE_LIGHT = "FFF4E5"
BLUE_LIGHT = "EAF2F8"
BORDER = "C9D1D6"

CLAIM_ORDER = ["C1", "C2", "C3", "C4"]
FAMILY_SUMMARIES = {
    "C1": {
        "Claim": "Mashiki 2026 service-presence trajectory",
        "Result": (
            "Positive passage trajectory observed; MMI is similar, PGA is lower, "
            "and the event outcomes are non-equivalent."
        ),
        "Measurement dependence": (
            "Yes—hazard ranking differs by metric; 2016 restrictions and 2026 "
            "probe passage measure different states."
        ),
        "Inference ceiling": (
            "No restoration-time ratio, causal effect, or claim of uniformly equal "
            "or stronger 2026 shaking."
        ),
    },
    "C2": {
        "Claim": "Yokoe faster restoration under stronger 2026 shaking",
        "Result": (
            "Faster in 2026 under both admissible endpoints; MMI and PGA both rank "
            "2026 shaking as stronger."
        ),
        "Measurement dependence": (
            "Direction is robust; endpoint meaning remains explicit and is not "
            "treated as damage equivalence."
        ),
        "Inference ceiling": (
            "No causal reconstruction effect or inference that shaking intensity "
            "equals bridge damage severity."
        ),
    },
    "C3": {
        "Claim": "Case-selection validity and contextual exclusions",
        "Result": (
            "Mashiki stages pass linkage at one verified corridor-cluster unit; "
            "Yokoe is the sole exact recurrence; contextual candidates fail core gates."
        ),
        "Measurement dependence": (
            "No for the current denominator audit; results must be rerun if official "
            "candidate universes expand."
        ),
        "Inference ceiling": (
            "No prevalence, selection-probability, or population inference; context "
            "cases do not enlarge the analytical core."
        ),
    },
    "C4": {
        "Claim": "Traceable source precedence",
        "Result": (
            "The official terminal record is retained through an explicit override "
            "of the lagging archive."
        ),
        "Measurement dependence": (
            "Source-rule dependent—the override and retrieval time must accompany "
            "the reported fact."
        ),
        "Inference ceiling": (
            "No unqualified claim that all official archives were synchronized at "
            "the observation time."
        ),
    },
}

EXPECTED_CLASSIFICATIONS = {
    "C1": {
        "robust for similar MMI",
        "measurement dependent",
        "robust positive observation",
        "non-comparable outcome",
    },
    "C2": {"direction robust"},
    "C3": {"audit robust", "unsupported for core"},
    "C4": {"source-rule robust"},
}


def checks_applied(subset: pd.DataFrame) -> str:
    labels: list[str] = []
    counts = subset["Sensitivity Dimension"].value_counts(sort=False)
    for dimension, count in counts.items():
        label = str(dimension).replace("MMI event comparability", "MMI comparability")
        label = label.replace("PGA event comparability", "PGA comparability")
        labels.append(f"{label} ({count})" if int(count) > 1 else label)
    return "; ".join(labels)


def build_table(data: pd.DataFrame) -> pd.DataFrame:
    if len(data) != 12 or set(data["Claim ID"]) != set(CLAIM_ORDER):
        raise RuntimeError("Unexpected claim-sensitivity matrix contract.")
    if data[["Result", "Classification", "Evidence", "Inference Ceiling"]].isna().any().any():
        raise RuntimeError("Every claim scenario must retain result and ceiling evidence.")

    rows: list[dict[str, str]] = []
    for claim_id in CLAIM_ORDER:
        subset = data.loc[data["Claim ID"].eq(claim_id)].copy()
        actual_classes = set(subset["Classification"])
        if actual_classes != EXPECTED_CLASSIFICATIONS[claim_id]:
            raise RuntimeError(f"Classification contract changed for {claim_id}.")
        summary = FAMILY_SUMMARIES[claim_id]
        rows.append(
            {
                "Claim": summary["Claim"],
                "Checks applied": checks_applied(subset),
                "Result": summary["Result"],
                "Measurement dependence": summary["Measurement dependence"],
                "Inference ceiling": summary["Inference ceiling"],
            }
        )
    return pd.DataFrame(rows)


def write_workbook(table: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Claim Boundaries"
    sheet.sheet_view.showGridLines = False

    header_row = 1
    for column_index, column_name in enumerate(table.columns, start=1):
        cell = sheet.cell(header_row, column_index, column_name)
        cell.font = Font(name="Aptos", size=9.4, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    thin = Side(style="thin", color=BORDER)
    body_start = header_row + 1
    body_end = body_start + len(table) - 1
    for row_index, record in enumerate(table.itertuples(index=False), start=body_start):
        for column_index, value in enumerate(record, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.font = Font(
                name="Aptos",
                size=8.9,
                bold=column_index == 1,
                color=TEXT,
            )
            if column_index == 1:
                fill = LIGHT_GRAY
            elif row_index % 2 == 1:
                fill = ORANGE_LIGHT
            else:
                fill = BLUE_LIGHT if column_index in {3, 4} else ROW_GRAY
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[row_index].height = 70

    widths = [35, 37, 52, 52, 52]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column_index)].width = width

    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A{header_row}:E{body_end}"
    excel_table = Table(
        displayName="ClaimRobustnessAndInferenceCeilings",
        ref=f"A{header_row}:E{body_end}",
    )
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.30
    sheet.page_margins.bottom = 0.30
    sheet.print_area = f"A1:E{body_end}"

    data_sheet = workbook.create_sheet("Data", 0)
    for column_index, column_name in enumerate(table.columns, start=1):
        data_sheet.cell(1, column_index, column_name)
    for row_index, record in enumerate(table.itertuples(index=False), start=2):
        for column_index, value in enumerate(record, start=1):
            data_sheet.cell(row_index, column_index, value)
    data_sheet.sheet_state = "hidden"
    workbook.active = workbook.sheetnames.index("Claim Boundaries")

    workbook.properties.title = "Claim Robustness and Inference Ceilings"
    workbook.properties.subject = "Kumamoto repeated-shock claim audit"
    workbook.properties.creator = "Mike Li"
    workbook.properties.description = (
        "Four claim families summarized from twelve prespecified sensitivity checks."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)


def main() -> None:
    data = pd.read_parquet(SOURCE)
    table = build_table(data)
    write_workbook(table)
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(f"Dimensions: {len(table)} rows x {len(table.columns)} columns")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
