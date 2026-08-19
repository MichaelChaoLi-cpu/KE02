#!/usr/bin/env python3
"""Candidate Selection Gate Summary.

Plan: Compress the 65-row candidate ledger into one auditable row for each
official denominator while preserving pathway-specific linkage rules.
Framework: AnaSOP Sections 5.3-5.5, Section 6.4, and Analytical Workflow
step 5. Counts describe audit coverage, not prevalence or selection odds.
"""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = (
    ROOT
    / "data/processed"
    / "repeated_shock_candidate_ledger_preprocessed.parquet"
)
OUT = (
    ROOT
    / "data/results/tables"
    / "Table_candidate_selection_gate_summary.xlsx"
)

NAVY = "243746"
WHITE = "FFFFFF"
TEXT = "263238"
LIGHT_GRAY = "E9EDF0"
ROW_GRAY = "F7F8F9"
ORANGE_LIGHT = "FFF4E5"
BLUE_LIGHT = "EAF2F8"
BORDER = "C9D1D6"

UNIVERSE_ORDER = [
    "post-2016 local road projects",
    "2016 officially restored bridges",
    "2026 bridge-damage records",
    "additional official reconstructions",
]
UNIVERSE_LABELS = {
    "post-2016 local road projects": "Post-2016 local road projects",
    "2016 officially restored bridges": "2016 officially restored bridges",
    "2026 bridge-damage records": "2026 bridge-damage records",
    "additional official reconstructions": "Additional official reconstructions",
}
EXCLUSION_LABELS = {
    "fail_new_post_2016_route_no_same_asset": "No same-asset 2016 failure",
    "fail_route_only_location_mismatch": "Route-only location mismatch",
    "no exact 2026 recurrence": "No exact 2026 recurrence",
    "fail_no_2016_rebuilt_bridge_on_route": "No rebuilt 2016 bridge on route",
    "fail_route_mismatch_spatial_proximity": "Route mismatch despite proximity",
    "exclude_hazard_and_geometry_not_comparable": (
        "Geometry mismatch and weaker 2026 shaking"
    ),
    "fail_initial_repeated_shock_screen": "Failed initial repeated-shock screen",
}


def valid_linkage(universe: str, subset: pd.DataFrame) -> tuple[int, str]:
    count = int((subset["Identity Gate"] & subset["Geometry Gate"]).sum())
    if universe == "post-2016 local road projects":
        return count, f"{count} corridor-cluster-linked stages"
    noun = "exact asset" if count == 1 else "geometry-compatible assets"
    return count, f"{count} {noun}"


def repeated_shock_evidence(subset: pd.DataFrame) -> tuple[int, str]:
    admissible_hazard = subset["Hazard Gate"].isin(["pass", "metric dependent"])
    count = int(
        (
            subset["Identity Gate"]
            & subset["Geometry Gate"]
            & subset["Outcome Gate"]
            & admissible_hazard
        ).sum()
    )
    if count == 0:
        return count, "0; 2026 shaking weaker"
    if subset.loc[admissible_hazard, "Hazard Gate"].eq("metric dependent").any():
        return count, f"{count}; hazard metric-dependent"
    return count, f"{count}; MMI and PGA pass"


def retained_role(universe: str, subset: pd.DataFrame) -> str:
    if universe == "post-2016 local road projects":
        count = int(
            subset["Core Status"].eq("include as one nested Mashiki cluster").sum()
        )
        return f"{count} stages → 1 Mashiki corridor"
    if universe in {
        "2016 officially restored bridges",
        "2026 bridge-damage records",
    }:
        count = int(
            subset["Core Status"].eq("include Yokoe Bridge core case").sum()
        )
        return f"{count} Yokoe audit record"
    count = int(subset["Core Status"].eq("contextual mechanism case only").sum())
    return f"{count} contextual; 0 core"


def dominant_exclusion(subset: pd.DataFrame) -> str:
    exclusions = subset.loc[subset["Exclusion Reason"].str.strip().ne(""), "Exclusion Reason"]
    if exclusions.empty:
        return "None"
    counts = exclusions.value_counts()
    code = str(counts.index[0])
    count = int(counts.iloc[0])
    if code not in EXCLUSION_LABELS:
        raise RuntimeError(f"Unmapped exclusion reason: {code}")
    return f"{EXCLUSION_LABELS[code]} (n={count})"


def build_table(data: pd.DataFrame) -> pd.DataFrame:
    if len(data) != 65 or set(data["Candidate Universe"]) != set(UNIVERSE_ORDER):
        raise RuntimeError("Unexpected candidate-ledger denominator contract.")
    if data["Candidate ID"].duplicated().any():
        raise RuntimeError("Candidate identifiers must be unique within the ledger.")
    included = data["Core Status"].str.startswith("include")
    if not data.loc[included, ["Identity Gate", "Geometry Gate", "Outcome Gate"]].all().all():
        raise RuntimeError("An included candidate fails a required Boolean gate.")
    if not data.loc[included, "Hazard Gate"].isin(["pass", "metric dependent"]).all():
        raise RuntimeError("An included candidate lacks admissible hazard evidence.")

    rows: list[dict[str, object]] = []
    for universe in UNIVERSE_ORDER:
        subset = data.loc[data["Candidate Universe"].eq(universe)].copy()
        linkage_count, linkage_label = valid_linkage(universe, subset)
        evidence_count, evidence_label = repeated_shock_evidence(subset)
        if evidence_count > linkage_count:
            raise RuntimeError(f"Repeated-shock evidence exceeds valid linkage: {universe}")
        rows.append(
            {
                "Candidate universe": UNIVERSE_LABELS[universe],
                "Records screened": int(len(subset)),
                "Valid linkage": linkage_label,
                "Repeated-shock evidence": evidence_label,
                "Retained role": retained_role(universe, subset),
                "Dominant exclusion": dominant_exclusion(subset),
            }
        )

    table = pd.DataFrame(rows)
    expected_screened = [12, 37, 13, 3]
    if table["Records screened"].tolist() != expected_screened:
        raise RuntimeError("Official denominator sizes changed unexpectedly.")
    return table


def write_workbook(table: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Selection Gates"
    sheet.sheet_view.showGridLines = False

    sheet.merge_cells("A1:F1")
    title = sheet["A1"]
    title.value = "Candidate Selection Gate Summary"
    title.font = Font(name="Aptos Display", size=15, bold=True, color=WHITE)
    title.fill = PatternFill("solid", fgColor=NAVY)
    title.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[1].height = 28

    header_row = 2
    for column_index, column_name in enumerate(table.columns, start=1):
        cell = sheet.cell(header_row, column_index, column_name)
        cell.font = Font(name="Aptos", size=9.2, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    sheet.row_dimensions[header_row].height = 30

    thin = Side(style="thin", color=BORDER)
    body_start = header_row + 1
    body_end = body_start + len(table) - 1
    for row_index, record in enumerate(table.itertuples(index=False), start=body_start):
        for column_index, value in enumerate(record, start=1):
            cell = sheet.cell(row_index, column_index, value)
            cell.font = Font(
                name="Aptos",
                size=9.2,
                bold=column_index == 1,
                color=TEXT,
            )
            if column_index == 1:
                fill = LIGHT_GRAY
            elif row_index % 2 == 1:
                fill = ORANGE_LIGHT
            else:
                fill = BLUE_LIGHT if column_index in {3, 4, 5} else ROW_GRAY
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(
                horizontal="center" if column_index == 2 else "left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin)
        sheet.row_dimensions[row_index].height = 42

    widths = [33, 16, 27, 30, 30, 38]
    for column_index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + column_index)].width = width

    sheet.freeze_panes = "B3"
    sheet.auto_filter.ref = f"A{header_row}:F{body_end}"
    excel_table = Table(
        displayName="CandidateSelectionGateSummary",
        ref=f"A{header_row}:F{body_end}",
    )
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.30
    sheet.page_margins.bottom = 0.30
    sheet.print_area = f"A1:F{body_end}"

    data_sheet = workbook.create_sheet("Data", 0)
    for column_index, column_name in enumerate(table.columns, start=1):
        data_sheet.cell(1, column_index, column_name)
    for row_index, record in enumerate(table.itertuples(index=False), start=2):
        for column_index, value in enumerate(record, start=1):
            data_sheet.cell(row_index, column_index, value)
    data_sheet.sheet_state = "hidden"
    workbook.active = workbook.sheetnames.index("Selection Gates")

    workbook.properties.title = "Candidate Selection Gate Summary"
    workbook.properties.subject = "Kumamoto repeated-shock case-selection audit"
    workbook.properties.creator = "Mike Li"
    workbook.properties.description = (
        "Four official candidate denominators summarized after identity, geometry, "
        "hazard, outcome, and role auditing."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)


def main() -> None:
    data = pd.read_parquet(SOURCE)
    table = build_table(data)
    write_workbook(table)
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(f"Dimensions: {len(table)} rows x {len(table.columns)} columns")
    print(table.to_string(index=False))


if __name__ == "__main__":
    main()
