#!/usr/bin/env python3
"""Two-Pathway Repeated-Shock Evidence Matrix.

Plan: Transpose the two-case evidence layer into a compact publication table
with one evidence dimension per row and one case per column.
Framework: AnaSOP Sections 5.1-5.5, Section 6.1, and Analytical Workflow
steps 1, 2, and 7. The cases remain deliberately non-pooled.
"""
from __future__ import annotations

import math
import re
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo


ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / "data/processed/repeated_shock_two_case_evidence_preprocessed.parquet"
OUT = (
    ROOT
    / "data/results/tables"
    / "Table_two_pathway_repeated_shock_evidence_matrix.xlsx"
)

NAVY = "243746"
WHITE = "FFFFFF"
TEXT = "263238"
LIGHT_GRAY = "E9EDF0"
ROW_GRAY = "F7F8F9"
ORANGE_LIGHT = "FFF4E5"
BLUE_LIGHT = "EAF2F8"
ORANGE = "D97706"
BLUE = "2563A8"
BORDER = "C9D1D6"


def compact_failure(case_id: str, text: str) -> str:
    if case_id == "MASHIKI-CENTRAL":
        full_count = text.lower().count("full closure")
        one_side_count = text.lower().count("one-side closure")
        if full_count != 2 or one_side_count != 1:
            raise RuntimeError("Unexpected Mashiki 2016 failure evidence.")
        return "2 subsegments under full closure; 1 under one-side closure."
    return text.rstrip(".") + "."


def compact_performance(case_id: str, text: str) -> str:
    if case_id == "MASHIKI-CENTRAL":
        hours = re.search(r"([0-9]+\.[0-9]+) hours", text)
        snapshots = re.search(r"(\d+/\d+) available snapshots", text)
        if hours is None or snapshots is None:
            raise RuntimeError("Unexpected Mashiki 2026 performance evidence.")
        return (
            f"Near-full probe passage first observed ≤{hours.group(1)} h; "
            f"passage in {snapshots.group(1)} available snapshots."
        )
    hours = re.search(r"([0-9]+\.[0-9]+) hours", text)
    if hours is None:
        raise RuntimeError("Unexpected Yokoe 2026 performance evidence.")
    days = float(hours.group(1)) / 24
    return f"General traffic restored {days:.2f} days after the event."


def format_date(value: object) -> str:
    timestamp = pd.Timestamp(value)
    return f"{timestamp.day} {timestamp.strftime('%b %Y')}"


def sentence_case(value: object, period: bool = False) -> str:
    text = str(value).strip()
    text = text[:1].upper() + text[1:]
    if period and not text.endswith("."):
        text += "."
    return text


def case_values(row: pd.Series) -> dict[str, str]:
    case_id = str(row["Case ID"])
    mmi_2016 = float(row["Mean MMI 2016"])
    mmi_2026 = float(row["Mean MMI 2026"])
    mmi_difference = float(row["MMI Difference (2026-2016)"])
    pga_2016 = float(row["Mean PGA 2016 (g)"])
    pga_2026 = float(row["Mean PGA 2026 (g)"])
    pga_ratio = float(row["PGA Ratio (2026/2016)"])
    return {
        "Adaptation pathway": sentence_case(row["Adaptation Pathway"]),
        "Documented intervention": sentence_case(row["Adaptation Description"]),
        "Adaptation completed": format_date(row["Adaptation Completion"]),
        "2016 failure evidence": compact_failure(
            case_id, str(row["Failure Evidence 2016"])
        ),
        "2026 observed performance": compact_performance(
            case_id, str(row["Performance Evidence 2026"])
        ),
        "Shaking: MMI": (
            f"{mmi_2016:.2f} → {mmi_2026:.2f} "
            f"(Δ {mmi_difference:+.2f})"
        ),
        "Shaking: PGA": (
            f"{pga_2016:.3f} → {pga_2026:.3f} g "
            f"(ratio {pga_ratio:.2f})"
        ),
        "Outcome comparability": sentence_case(row["Outcome Comparability"]),
        "Evidence level": sentence_case(row["Evidence Level"]),
        "Permissible claim": sentence_case(row["Permissible Claim"], period=True),
        "Inference ceiling": sentence_case(row["Inference Ceiling"], period=True),
    }


def build_table(data: pd.DataFrame) -> pd.DataFrame:
    expected_cases = {"MASHIKI-CENTRAL", "YOKOE-BRIDGE"}
    if len(data) != 2 or set(data["Case ID"]) != expected_cases:
        raise RuntimeError("The evidence layer must contain the two declared cases.")
    if not data["Core Case"].all():
        raise RuntimeError("A declared evidence-matrix case is not marked core.")

    indexed = data.set_index("Case ID", drop=False)
    mashiki = case_values(indexed.loc["MASHIKI-CENTRAL"])
    yokoe = case_values(indexed.loc["YOKOE-BRIDGE"])
    dimensions = list(mashiki)
    if dimensions != list(yokoe) or len(dimensions) != 11:
        raise RuntimeError("Unexpected evidence-matrix dimension contract.")
    return pd.DataFrame(
        {
            "Evidence dimension": dimensions,
            "Mashiki Central Line": [mashiki[value] for value in dimensions],
            "Yokoe Bridge": [yokoe[value] for value in dimensions],
        }
    )


def write_workbook(table: pd.DataFrame) -> None:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Evidence Matrix"
    sheet.sheet_view.showGridLines = False

    header_row = 1
    for column_index, column_name in enumerate(table.columns, start=1):
        cell = sheet.cell(header_row, column_index, column_name)
        cell.font = Font(name="Aptos", size=10, bold=True, color=WHITE)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="left", vertical="center")
    sheet.row_dimensions[header_row].height = 25

    thin = Side(style="thin", color=BORDER)
    body_start = header_row + 1
    for row_offset, record in enumerate(table.itertuples(index=False), start=body_start):
        values = list(record)
        maximum_lines = 1
        for column_index, value in enumerate(values, start=1):
            cell = sheet.cell(row_offset, column_index, value)
            cell.font = Font(
                name="Aptos",
                size=9.2,
                bold=column_index == 1,
                color=TEXT,
            )
            if column_index == 1:
                fill = LIGHT_GRAY
            elif row_offset % 2 == 0:
                fill = ORANGE_LIGHT if column_index == 2 else BLUE_LIGHT
            else:
                fill = ROW_GRAY
            cell.fill = PatternFill("solid", fgColor=fill)
            cell.alignment = Alignment(
                horizontal="left",
                vertical="top",
                wrap_text=True,
            )
            cell.border = Border(bottom=thin)
            estimated_width = 28 if column_index == 1 else 55
            maximum_lines = max(
                maximum_lines,
                math.ceil(len(str(value)) / estimated_width),
            )
        sheet.row_dimensions[row_offset].height = max(24, 14 * maximum_lines + 8)

    sheet.column_dimensions["A"].width = 28
    sheet.column_dimensions["B"].width = 55
    sheet.column_dimensions["C"].width = 55
    sheet.freeze_panes = "B2"
    sheet.auto_filter.ref = f"A{header_row}:C{body_start + len(table) - 1}"

    excel_table = Table(
        displayName="RepeatedShockEvidenceMatrix",
        ref=f"A{header_row}:C{body_start + len(table) - 1}",
    )
    excel_table.tableStyleInfo = TableStyleInfo(
        name="TableStyleMedium2",
        showFirstColumn=False,
        showLastColumn=False,
        showRowStripes=False,
        showColumnStripes=False,
    )
    sheet.add_table(excel_table)

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_options.horizontalCentered = True
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.page_margins.left = 0.25
    sheet.page_margins.right = 0.25
    sheet.page_margins.top = 0.30
    sheet.page_margins.bottom = 0.30
    sheet.print_area = f"A1:C{body_start + len(table) - 1}"

    data_sheet = workbook.create_sheet("Data", 0)
    for column_index, column_name in enumerate(table.columns, start=1):
        data_sheet.cell(1, column_index, column_name)
    for row_index, record in enumerate(table.itertuples(index=False), start=2):
        for column_index, value in enumerate(record, start=1):
            data_sheet.cell(row_index, column_index, value)
    data_sheet.sheet_state = "hidden"
    workbook.active = workbook.sheetnames.index("Evidence Matrix")

    workbook.properties.title = "Two-Pathway Repeated-Shock Evidence Matrix"
    workbook.properties.subject = "Kumamoto repeated-shock mechanism evidence"
    workbook.properties.creator = "Mike Li"
    workbook.properties.description = (
        "Compact transposed evidence matrix for two deliberately non-pooled cases."
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(OUT)


def main() -> None:
    data = pd.read_parquet(SOURCE)
    table = build_table(data)
    write_workbook(table)
    print(f"Saved: {OUT.relative_to(ROOT)}")
    print(f"Dimensions: {len(table)} rows x {len(table.columns)} columns")


if __name__ == "__main__":
    main()
